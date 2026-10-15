def find_shape(slide, shape_name):

    for shape in slide.shapes:

        if shape.name == shape_name:
            return shape

    raise ValueError(
        f"{shape_name} not found"
    )

from pptx import Presentation

from ppt_utils import find_shape

prs = Presentation(
    "input/novaretail_template.pptx"
)

slide = prs.slides[0]

shape = find_shape(
    slide,
    "deck_title"
)

print(shape.text)