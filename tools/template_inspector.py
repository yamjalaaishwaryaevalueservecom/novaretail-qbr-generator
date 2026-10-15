from pptx import Presentation

prs = Presentation(
    "input/novaretail_template.pptx"
)

for slide_number, slide in enumerate(
    prs.slides,
    start=1
):
    print(f"\n===== SLIDE {slide_number} =====")

    for shape in slide.shapes:

        if not shape.has_text_frame:
            continue

        for paragraph in shape.text_frame.paragraphs:

            for run in paragraph.runs:
                print(
                    f"Shape: {shape.name} | "
                    f"Font: {run.font.name} | "
                    f"Size: {run.font.size} | "
                    f"Bold: {run.font.bold}"
                )